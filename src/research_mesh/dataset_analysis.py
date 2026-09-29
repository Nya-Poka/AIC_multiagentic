from __future__ import annotations

import csv
import json
import math
import statistics
from collections import Counter
from pathlib import Path
from typing import Any

from .artifacts import ArtifactError, LocalArtifactStore
from .schemas import DatasetAnalysisSpec, DatasetArtifactRef


class DatasetAnalysisError(ValueError):
    """The referenced dataset cannot be analyzed safely."""


_MISSING = {"", "na", "n/a", "null", "none", "nan", "."}


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel
        return [dict(row) for row in csv.DictReader(handle, dialect=dialect)]


def _read_json(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(payload, dict):
        payload = payload.get("records", payload.get("data", payload))
    if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
        raise DatasetAnalysisError("JSON dataset must be an array of objects")
    return [dict(item) for item in payload]


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on optional runtime extra
        raise DatasetAnalysisError(
            "XLSX analysis requires the optional openpyxl dependency"
        ) from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = [
                str(value).strip() if value is not None else "" for value in next(rows)
            ]
        except StopIteration:
            return []
        if not all(headers) or len(set(headers)) != len(headers):
            raise DatasetAnalysisError(
                "XLSX first row must contain unique non-empty headers"
            )
        return [dict(zip(headers, values, strict=False)) for values in rows]
    finally:
        if "workbook" in locals():
            workbook.close()


def _load_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _read_csv(path)
    elif suffix == ".json":
        rows = _read_json(path)
    elif suffix == ".xlsx":
        rows = _read_xlsx(path)
    else:
        raise DatasetAnalysisError("unsupported dataset extension")
    if not rows:
        raise DatasetAnalysisError("dataset contains no records")
    if len(rows) > 100_000:
        raise DatasetAnalysisError("dataset exceeds the 100000-row analysis limit")
    return rows


def _missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip().casefold() in _MISSING)


def _number(value: Any) -> float | None:
    if _missing(value) or isinstance(value, bool):
        return None
    try:
        parsed = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return parsed if math.isfinite(parsed) else None


def _quantile(values: list[float], q: float) -> float:
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    position = (len(ordered) - 1) * q
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def _numeric_summary(values: list[float]) -> dict[str, float | int | None]:
    return {
        "count": len(values),
        "mean": round(statistics.fmean(values), 6),
        "std": round(statistics.stdev(values), 6) if len(values) > 1 else None,
        "min": round(min(values), 6),
        "q1": round(_quantile(values, 0.25), 6),
        "median": round(statistics.median(values), 6),
        "q3": round(_quantile(values, 0.75), 6),
        "max": round(max(values), 6),
    }


def _pearson(pairs: list[tuple[float, float]]) -> dict[str, Any] | None:
    if len(pairs) < 4:
        return None
    xs = [item[0] for item in pairs]
    ys = [item[1] for item in pairs]
    if len(set(xs)) < 2 or len(set(ys)) < 2:
        return None
    x_mean = statistics.fmean(xs)
    y_mean = statistics.fmean(ys)
    numerator = sum((x - x_mean) * (y - y_mean) for x, y in pairs)
    denominator = math.sqrt(
        sum((x - x_mean) ** 2 for x in xs) * sum((y - y_mean) ** 2 for y in ys)
    )
    if denominator == 0:
        return None
    r = max(-1.0, min(1.0, numerator / denominator))
    if len(pairs) > 3 and abs(r) < 1:
        z = math.atanh(r)
        margin = 1.96 / math.sqrt(len(pairs) - 3)
        lower, upper = math.tanh(z - margin), math.tanh(z + margin)
    else:
        lower, upper = r, r
    return {
        "n": len(pairs),
        "pearson_r": round(r, 6),
        "confidence_interval_95": [round(lower, 6), round(upper, 6)],
        "method": "pearson-correlation-with-fisher-z-ci",
    }


def analyze_dataset(
    reference: DatasetArtifactRef,
    spec: DatasetAnalysisSpec | None,
    *,
    store: LocalArtifactStore | None = None,
) -> dict[str, Any]:
    try:
        stored = (store or LocalArtifactStore()).resolve_dataset(reference)
        rows = _load_rows(stored.path)
    except (ArtifactError, OSError, UnicodeError, json.JSONDecodeError, csv.Error) as exc:
        raise DatasetAnalysisError(str(exc)) from exc

    columns = list(dict.fromkeys(key for row in rows for key in row))
    if not columns:
        raise DatasetAnalysisError("dataset contains no columns")
    column_profiles: dict[str, dict[str, Any]] = {}
    numeric_values: dict[str, list[float]] = {}
    for column in columns:
        raw = [row.get(column) for row in rows]
        missing_count = sum(_missing(value) for value in raw)
        numbers = [number for value in raw if (number := _number(value)) is not None]
        non_missing_count = len(raw) - missing_count
        numeric = non_missing_count > 0 and len(numbers) / non_missing_count >= 0.9
        profile: dict[str, Any] = {
            "type": "numeric" if numeric else "categorical",
            "non_missing": non_missing_count,
            "missing": missing_count,
            "missing_rate": round(missing_count / len(rows), 6),
            "unique": len({str(value) for value in raw if not _missing(value)}),
        }
        if numeric:
            numeric_values[column] = numbers
            profile["summary"] = _numeric_summary(numbers)
        else:
            profile["top_values"] = [
                {"value": value, "count": count}
                for value, count in Counter(
                    str(value) for value in raw if not _missing(value)
                ).most_common(5)
            ]
        column_profiles[column] = profile

    resolved_spec = spec or DatasetAnalysisSpec()
    correlations: list[dict[str, Any]] = []
    outcome = resolved_spec.outcome
    if outcome:
        if outcome not in columns:
            raise DatasetAnalysisError(f"outcome column not found: {outcome}")
        if outcome not in numeric_values:
            raise DatasetAnalysisError("outcome must be numeric for the current analysis engine")
        for exposure in resolved_spec.exposures:
            if exposure not in columns:
                raise DatasetAnalysisError(f"exposure column not found: {exposure}")
            pairs = [
                (x, y)
                for row in rows
                if (x := _number(row.get(exposure))) is not None
                and (y := _number(row.get(outcome))) is not None
            ]
            result = _pearson(pairs)
            correlations.append(
                {
                    "exposure": exposure,
                    "outcome": outcome,
                    "result": result,
                    "warning": (
                        None
                        if result is not None
                        else "insufficient complete numeric pairs or zero variance"
                    ),
                }
            )

    return {
        "artifact": reference.model_dump(mode="json"),
        "row_count": len(rows),
        "column_count": len(columns),
        "columns": column_profiles,
        "analysis_spec": resolved_spec.model_dump(mode="json"),
        "associations": correlations,
        "warnings": [
            "Associations are descriptive and do not establish causality.",
            "The deterministic engine does not currently adjust for covariates; listed covariates are preserved for review.",
        ],
        "reproducibility": {
            "engine": "research-mesh-dataset-audit-v1",
            "input_sha256": reference.sha256,
            "network_access": False,
            "raw_rows_returned": False,
        },
    }
